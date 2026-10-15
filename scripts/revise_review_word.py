"""Clean review wording and create an editable review table."""
from pathlib import Path
import ast, json, shutil, subprocess, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'scripts/review_originals.py'
s=p.read_text(encoding='utf-8')
archive=ROOT/'docs/archive/before-word-edit';archive.mkdir(parents=True,exist_ok=True)
for src in [p,ROOT/'docs/part3-review.md',ROOT/'docs/review-data.json']:
    if not (archive/src.name).exists(): shutil.copy2(src,archive/src.name)
old=json.loads((ROOT/'docs/review-data.json').read_text(encoding='utf-8'))
issues=[
'左侧背景和人物背后的布面偏暗，几处阴影连成一片，材质的差别不够清楚。',
'手背的高光有些抢眼，指尖挤在眼睛附近，让脸部周围显得局促；皮肤也略显光滑，少了些纹理。',
'头巾下方的阴影偏重，几道褶皱混在一起；蓝眼睛和青绿色头巾的颜色都较浓，略显刻意。',
'颈部珠串和肩部金纹过于密集，亮点容易抢走对脸部的注意，首饰之间的轮廓也不够清楚。',
'上方的大泡泡与手臂重叠，背景人群又比较密集；两侧人物被画面边缘截断，取景稍显拥挤。',
'左侧大招牌和道路上方几块小招牌的字形含混，仔细看时比较出戏；近处人物和路面细节也偏简单。',
'人物周围的汽车和植被太暗，部分轮廓融进背景。气氛很强，但暗处能读到的环境细节偏少。',
'太阳附近亮得有些发白，云和树枝的细节被盖住；湖水与树叶的颜色偏浓，右侧密林也略显杂乱。',
'衣服和伞面都有密集花纹，腰部与手臂的轮廓容易混在一起；座椅旁的白色织物过亮，褶皱不够清楚。',
'两侧花束和烛光都很醒目，分散了对香水瓶的注意。瓶子周围留白较少，作为商品图，主角还不够突出。',
'白猫胸口的亮部接近一片白，毛发层次偏少；画面下沿截去了部分抱猫动作，人与猫的姿态不够完整。',
'银色手包与裙子的亮部颜色接近，包的边缘不够突出；站姿和背景都比较常规，画面缺少让人记住的细节。',
'中央深色物件与后方黑瓶融在一起，轮廓不太清楚；高脚盘底座附近堆了几件小物，让桌面显得有些挤。',
'两侧金色粒子太亮，容易抢过香水瓶和蜡烛；瓶身下半部的闪点过密，玻璃的通透感反而弱了。',
'右侧银色袖扣的高光连成亮斑，中央黑色袖扣的纹理偏暗；物件都集中在下方，上半部留白稍多。',
'花瓶、壶盖和牛角包同时悬空，但画面没有进一步展开这种失重感，与周围安静的日常陈设有些脱节。',
'右侧人物的脸与车辆反射叠在一起，表情不容易看清；底部招牌文字被截断，画面边缘略显仓促。',
'两侧灯牌比人物的脸更亮，前景装饰又有很多零碎高光，视线容易被带走，人物的存在感偏弱。',
'右上方的裸地和水带较亮，稍微分散了对人物的注意；草地绿色偏浓，局部看起来有些鲜艳。',
'近处两根电杆占据较大面积，把道路切得有些零碎；远处招牌和灯光较密，街道尽头稍显杂乱。',
]
tree=ast.parse(s)
assignment=next(n for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='reviews' for t in n.targets))
entries=ast.literal_eval(assignment.value)
for entry,issue in zip(entries,issues): s=s.replace(repr(entry[4]),repr(issue),1)
start=s.index('2026-09-29逐张查看');end=s.index('## 2. 评分标准与权重',start)
s=s[:start]+'图片编号对应图集01至20。先看整体观感，再检查局部细节，按构图、光影、结构、表达和完成度五个维度评分。\n\n'+s[end:]
s=s.replace('以当前尺寸观看时的读取成本、裁切及需要修复的范围；未知授权不计分','主体是否清楚、裁切是否合适，以及局部调整的工作量')
start=s.index('严重性规则：');end=s.index('总分含义：',start)
s=s[:start]+s[end:]
s=s.replace('暗部与文件分辨率限制使其排在19之后','背景暗部层次略弱，因此排在19之后')
s=s.replace('11也为4.65，本次优先07的环境、尺度与光线共同组织，属于同分后的明确编辑取舍，不声称两者存在客观显著差距。','11也为4.65；07通过光束和人物尺度营造的悬念更鲜明，因此排在第三。')
s=s.replace('它们不是不可用，排序依据是修正问题的范围与不确定性，不强行降为1或2分。','三张都有明显亮点，但招牌细节、场景表达和视觉主次分别拖累了整体效果。')
s=s.replace('若明确用于纯符号化世界观，可减轻文字相关扣分。','')
s=s.replace('花瓶、壶盖和牛角包的悬浮可直接看到，但缺少提示词时无法确定意图。应先明确是否做失重场景，再统一运动/支撑逻辑或修复落地关系；比局部亮度调整涉及更大的表达决策。','花瓶、壶盖和牛角包的悬浮很醒目，但失重感与日常陈设之间缺少呼应。需要重新组织场景，让这种异常成为明确的画面主题。')
start=s.index('## 6. 关键修订与一致性说明');end=s.index("'''",start)
s=s[:start]+s[end:]
p.write_text(s,encoding='utf-8')
subprocess.run([sys.executable,str(p)],check=True)
records=json.loads((ROOT/'docs/review-data.json').read_text(encoding='utf-8'))
assert [(r['id'],r['score'],r['dimensions']) for r in old]==[(r['id'],r['score'],r['dimensions']) for r in records]
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(29.7);sec.page_height=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(1.5);sec.left_margin=sec.right_margin=Cm(1.6)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(3)
doc.styles['Title'].font.size=Pt(22)
widths=[1.1,1.1,8.5,8.0,7.8]
for group in range(4):
    if group:doc.add_page_break()
    doc.add_paragraph('20张图片美学测评' if group==0 else '20张图片美学测评（续）',style='Title')
    if group==0:doc.add_paragraph('5分：优秀　4分：整体好，局部可修　3分：有亮点，问题明显　2分：主要问题影响观看　1分：不可用')
    table=doc.add_table(rows=1,cols=5);table.autofit=False
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    for cell,label in zip(table.rows[0].cells,['编号','总分','核心判断依据','优点','问题']):cell.text=label
    for r in records[group*5:group*5+5]:
        cells=table.add_row().cells
        for cell,text in zip(cells,[r['id'],str(r['score'])+'分',r['basis'],r['pros'],r['issues']]):cell.text=text
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell,w in zip(row.cells,widths):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'24475B' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.12
                for run in para.runs:
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
out=ROOT/'output/docx/20张图片美学测评-可编辑版.docx';out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
assert sum(len(t.rows)-1 for t in Document(out).tables)==20
print(out)
