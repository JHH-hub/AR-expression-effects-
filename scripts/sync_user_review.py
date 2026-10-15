from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Pt,RGBColor
import ast,json,shutil,subprocess,sys
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'output/docx/20张图片美学测评-可编辑版.docx'
backup=ROOT/'docs/archive'/('user-edit-'+datetime.now().strftime('%Y%m%d-%H%M%S'))
backup.mkdir(parents=True);shutil.copy2(p,backup/p.name)
d=Document(p)
rows={row.cells[0].text:row for t in d.tables for row in t.rows[1:] if row.cells[0].text.isdigit()}
def put(n,col,text):
    c=rows[n].cells[col];c.paragraphs[0].clear().add_run(text)
put('01',3,rows['01'].cells[3].text.replace('记实','纪实'))
put('13',4,rows['13'].cells[4].text.replace('带子','袋子'))
put('11',4,'无明显问题。')
put('18',2,'男孩居中，平静表情与周围热闹的派对装饰形成反差；右侧地面延伸到身体旁，但躯干位置过低，画面没有交代完整的坐姿或蹲姿。')
put('18',3,'冷暖灯光和彩色反射营造出浓烈的派对气氛，脸部仍可辨识，颗粒质感鲜明。')
put('18',4,'男孩的躯干与右侧地面的高度关系不协调，下半身像陷入地面，人物姿态不成立。前景彩球、杯子和零碎高光过密，画面也显得杂乱。')
values={n:[c.text for c in row.cells] for n,row in rows.items()}
assert len(values)==20
source=ROOT/'scripts/review_originals.py';s=source.read_text(encoding='utf-8')
shutil.copy2(source,backup/source.name)
node=next(n for n in ast.parse(s).body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='reviews' for t in n.targets))
reviews=[(int(n),None,*v[2:]) for n,v in sorted(values.items())]
lines=s.splitlines(keepends=True)
s=''.join(lines[:node.lineno-1])+'reviews='+repr(reviews)+'\n'+'review_scores='+repr({int(n):int(v[1].replace('分','')) for n,v in values.items()})+'\n'+''.join(lines[node.end_lineno:])
s=s.replace('raw=round(sum(a*b for a,b in zip(dims,weights)),2)','raw=None').replace('score=math.floor(raw+.5)','score=review_scores[n]')
a=s.index('最看重构图与结构');b=s.index('总分含义：',a)
s=s[:a]+'最看重构图与结构，两者合计50%。权重用于确定评价重点，总分结合问题的严重程度判断。局部色彩、裁切或层次问题主要影响4分与5分的区别；主体的人体结构或空间关系明显失常，会影响整张图的可用性。\n\n'+s[b:]
a=s.index("audit='''");b=s.index("(ROOT/'docs/part3-review.md').write_text",a)
ranking='''

## 4. 最好3张与最差3张

最好顺序：19、07、11，均为5分。

- 第一名19：草坡曲线、长影和人物的位置相互呼应，侧光把地形层次表现得很清楚。画面元素少，构图却很完整。
- 第二名07：斜向光束把视线引向小人物，普通住宅与异常光照形成悬念，明暗对比有力度。相比19，电线等环境细节略显杂乱。
- 第三名11：红衣、白猫、黄门与紫墙形成鲜明的色块关系，衣料和猫毛的质感对比也很突出。03同为5分，但11的整体构成更鲜明，因此排在前三。

最差顺序：18、02、06，分别为1分、1分、2分。

- 第一位18：躯干与地面的高度关系不协调，下半身像陷入地面，直接影响主体姿态的可信度；密集装饰和高光进一步加重了画面的杂乱。
- 第二位02：手部多指问题位于脸部附近，正好处在视觉焦点，影响人像的可用性。相比18，问题集中在手部，修正范围较小。
- 第三位06：多块招牌的字形含混，近处人物与路面细节也偏简单。夜景气氛成立，但经不起细看。
'''
s=s[:a]+"audit=''\nranking="+repr(ranking)+'\n'+s[b:]
s=s.replace("['编号','原文件','C','L','S','E','U','加权分','总分','原评分','核心判断依据','优点','问题']","['编号','原文件','总分','原评分','核心判断依据','优点','问题']")
s=s.replace("[r['id'],r['file'],*r['dimensions'],r['raw'],r['score'],r['old_score'],r['basis'],r['pros'],r['issues']]","[r['id'],r['file'],r['score'],r['old_score'],r['basis'],r['pros'],r['issues']]")
source.write_text(s,encoding='utf-8')
d.add_page_break();d.add_paragraph('评分标准与排序',style='Title')
d.add_paragraph('构图与视觉层级25%　光影与色彩20%　结构与细节25%　风格与表达20%　成图可用性10%')
d.add_paragraph('权重用于确定评价重点，总分结合问题的严重程度判断。主体的人体结构或空间关系明显失常，会影响整张图的可用性。')
for line in ranking.splitlines():
    if not line or line.startswith('##'):continue
    para=d.add_paragraph(line.removeprefix('- '))
    if line.startswith(('最好顺序','最差顺序')):
        for r in para.runs:r.bold=True
for name in ['Title','Heading 1','Heading 2']:d.styles[name].font.color.rgb=RGBColor(0,0,0)
for root in [d.styles.element,d.element]:
    for el in list(root.xpath('.//w:pBdr')):el.getparent().remove(el)
d.save(p.with_name('20张图片美学测评-更新版.docx'))
subprocess.run([sys.executable,str(source)],check=True)
print('Saved editable Word and synchronized 20 user scores; removed obsolete calculated subscores.')
