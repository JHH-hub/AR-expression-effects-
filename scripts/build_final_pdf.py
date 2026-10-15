"""Compile the three exam parts, importing the user's latest Word edits."""
from pathlib import Path
from docx import Document
import json,re
ROOT=Path(__file__).resolve().parents[1]
# Reuse layout helpers without regenerating earlier review artifacts.
helper=(ROOT/'scripts/build_submission.py').read_text(encoding='utf-8').split("parts=[(ROOT/")[0]
exec(compile(helper,str(ROOT/'scripts/build_submission.py'),'exec'))
word=ROOT/'output/docx/20张图片美学测评-更新版.docx'
d=Document(word)
current=[]
for t in d.tables:
    for row in t.rows[1:]:
        cells=[c.text.strip() for c in row.cells]
        if len(cells)==5 and cells[0].isdigit():current.append(cells)
assert [r[0] for r in current]==[f'{n:02}' for n in range(1,21)]
scores={r[0]:int(r[1].replace('分','')) for r in current}
for r in review_records:r['score']=scores[r['id']]
paragraphs=[p.text.strip() for p in d.paragraphs if p.text.strip()]
ranking=paragraphs[paragraphs.index('最好顺序：19、07、11，均为5分。'):]
part3=(ROOT/'docs/part3-review.md').read_text(encoding='utf-8').split('## 3. 20张图片测评表')[0]
part3+='## 3. 20张图片测评表\n\n| 图片编号 | 总分 | 核心判断依据 | 优点 | 问题 |\n| --- | --- | --- | --- | --- |\n'
part3+='\n'.join('| '+' | '.join(row)+' |' for row in current)
part3+='\n\n## 4. 最好3张与最差3张\n\n'+'\n\n'.join(ranking)+'\n'
parts=[(ROOT/'docs/part1-pipeline.md').read_text(encoding='utf-8'),(ROOT/'docs/part2-demo.md').read_text(encoding='utf-8'),part3]
parts[1]=parts[1].replace('具体记录见 verification.md，未测项不写为通过。','')
recap=re.search(r'## 3\. Vibecoding.*?\n\n(.*?)\n\n## 4',parts[1],re.S).group(1)
assert len(re.sub(r'\s','',recap))<=200
flow=[Spacer(1,72),p('AI PRODUCT DESIGN','muted'),p('AI产品设计测试题','cover'),Spacer(1,16),p('自动化生产管线 · AR互动原型 · 图像美学评测','h2'),Spacer(1,45),p('内容目录','h2'),p('Part 1　全球50大区礼物资产生产管线'),p('Part 2　微笑落雨，大笑烟花'),p('Part 3　20张图片美学测评'),Spacer(1,32)]
flow+=pdf_blocks('Live Demo：https://facefx.pxlsan.cn/\n\nPublic Source：https://github.com/JHH-hub/AR-expression-effects-')
for md in parts:flow.append(PageBreak());flow.extend(pdf_blocks(md))
flow.extend(thumbnail_pages())
def footer(c,doc):
    c.setFont('CN',8);c.setFillColor(colors.HexColor('#617183'))
    c.drawString(44,A4[1]-29,'AI 产品设计测试题')
    c.drawRightString(A4[0]-44,23,str(doc.page))
target=OUT/'pdf/AI产品设计测试题-完整作答.pdf'
SimpleDocTemplate(str(target),pagesize=A4,rightMargin=44,leftMargin=44,topMargin=49,bottomMargin=49,title='AI产品设计测试题',author='').build(flow,onFirstPage=footer,onLaterPages=footer)
(OUT/'AI产品设计测试题-完整作答.md').write_text('\n\n---\n\n'.join(parts),encoding='utf-8')
reader=PdfReader(target);text='\n'.join(x.extract_text() or '' for x in reader.pages)
links=[]
for page in reader.pages:
    for ref in page.get('/Annots',[]):
        a=ref.get_object().get('/A',{});uri=a.get('/URI')
        if uri:links.append(uri)
assert 'https://facefx.pxlsan.cn/' in links and 'https://github.com/JHH-hub/AR-expression-effects-' in links
assert 'Vogue' in text and '较为粗糙' in text and '下半身像陷入地面' in text
assert all('Part '+str(n) in text for n in [1,2,3])
print(json.dumps({'path':str(target),'pages':len(reader.pages),'bytes':target.stat().st_size,'links':links,'recap_chars':len(re.sub(r'\s','',recap))},ensure_ascii=False))
